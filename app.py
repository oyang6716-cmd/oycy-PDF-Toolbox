"""oycy PDF 工具箱：Windows 离线桌面版，Tk GUI。"""
import os, sys, json, io, re, glob, shutil, tempfile, threading, subprocess, urllib.request, urllib.error, zipfile, ctypes
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import fitz
from PIL import Image, ImageDraw, ImageFont
import keyring

APP_NAME = 'oycy的PDF工具箱'
SERVICE = 'oycy-pdf-ai'
CONFIG = Path(os.getenv('APPDATA', str(Path.home()))) / 'oycy-PDF-Toolbox' / 'settings.json'
DEFAULTS = dict(api_base='https://api.openai.com/v1', model='gpt-4.1-mini', quality='标准', output_dir='', remember=True)
BLUE, RED, YELLOW, INK, PAPER = '#214ccf', '#e5412d', '#f2c33e', '#141414', '#fffaf0'


def enable_high_dpi(root=None):
    """Enable sharper rendering on Windows and set a sensible Tk scaling."""
    if sys.platform.startswith('win'):
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            try:
                ctypes.windll.user32.SetProcessDPIAware()
            except Exception:
                pass
    if root is not None:
        try:
            current = float(root.tk.call('tk', 'scaling'))
        except Exception:
            current = 1.0
        try:
            px_per_inch = root.winfo_fpixels('1i')
            recommended = max(1.35, min(2.0, px_per_inch / 72.0))
        except Exception:
            recommended = 1.35
        if current < recommended:
            try:
                root.tk.call('tk', 'scaling', recommended)
            except Exception:
                pass


def load_cfg():
    try:
        return {**DEFAULTS, **json.loads(CONFIG.read_text(encoding='utf-8'))}
    except Exception:
        return DEFAULTS.copy()


def save_cfg(d):
    CONFIG.parent.mkdir(parents=True, exist_ok=True)
    CONFIG.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding='utf-8')


def points(spec, max_page):
    if not spec.strip(): return list(range(max_page))
    out = []
    for token in re.split(r'[,，\s]+', spec.strip()):
        if not token: continue
        if '-' in token:
            a, b = token.split('-', 1)
            a, b = int(a), int(b)
            if not (1 <= a <= b <= max_page): raise ValueError(f'页面范围超出有效范围：{token}')
            out.extend(range(a - 1, b))
        else:
            n = int(token)
            if not 1 <= n <= max_page: raise ValueError(f'页码超出范围：{n}')
            out.append(n - 1)
    return out


def draw_text_overlay(doc, text, kind):
    # 使用 PIL 字体渲染为图片后嵌入 PDF；更可靠地覆盖中文字符。
    fontpaths = [r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\simhei.ttf',
                 '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    fontpath = next((p for p in fontpaths if Path(p).exists()), None)
    for i, page in enumerate(doc):
        p = page.rect
        if kind == '页码':
            s = f'{i+1} / {len(doc)}'; fs = 24; x, y = p.width / 2, p.height - 31
        else:
            s = text; fs = 48; x, y = p.width / 2, p.height / 2
        font = ImageFont.truetype(fontpath, fs) if fontpath else ImageFont.load_default()
        im = Image.new('RGBA', (max(600, int(p.width)), 140), (255,255,255,0))
        dd = ImageDraw.Draw(im)
        dd.text((im.width/2, 70), s, fill=(220,40,30,230), font=font, anchor='mm')
        with io.BytesIO() as buf:
            im.save(buf, format='PNG')
            w = min(p.width*0.9, im.width)
            h = w * im.height/im.width
            rect = fitz.Rect(x-w/2, y-h/2, x+w/2, y+h/2)
            page.insert_image(rect, stream=buf.getvalue(), overlay=True)


def run_operation(operation, inputs, dest, quality, page_spec='', extra=''):
    if not inputs: raise ValueError('请先选择文件。')
    q = {'轻量': (90, 0.85), '标准': (130, 0.88), '高清': (180, 0.93), '超清': (240, 0.97)}[quality]
    dpi, jpgquality = q
    dest = Path(dest); dest.parent.mkdir(parents=True, exist_ok=True)
    op = operation
    if op == '合并 PDF':
        merged = fitz.open()
        try:
            for f in inputs:
                with fitz.open(f) as d:
                    if d.needs_pass: raise ValueError('包含加密 PDF，请先解锁。')
                    merged.insert_pdf(d)
            merged.save(str(dest), garbage=4, deflate=True)
        finally: merged.close()
        return str(dest)
    if op == '图片转 PDF':
        d = fitz.open()
        try:
            for f in inputs:
                with Image.open(f) as im:
                    if im.mode not in ('RGB', 'L'): im = im.convert('RGB')
                    with io.BytesIO() as tmp:
                        im.save(tmp, 'JPEG', quality=round(jpgquality*100), optimize=True)
                        data = tmp.getvalue()
                pix = fitz.open(stream=data, filetype='jpeg')
                pdf = fitz.open('pdf', pix.convert_to_pdf())
                d.insert_pdf(pdf)
                pix.close();pdf.close()
            d.save(str(dest), garbage=4, deflate=True)
        finally:d.close()
        return str(dest)
    if op == 'Word 转 PDF':
        exe = shutil.which('soffice') or shutil.which('libreoffice')
        if not exe: raise RuntimeError('此功能需要另外安装 LibreOffice（免费），并确保 soffice 可在 PATH 中找到。')
        if len(inputs) != 1: raise ValueError('每次请选择一个 Word 文件。')
        with tempfile.TemporaryDirectory() as td:
            r = subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', td, inputs[0]],
                               capture_output=True, text=True, timeout=120)
            files = list(Path(td).glob('*.pdf'))
            if r.returncode or not files: raise RuntimeError(r.stderr or r.stdout or 'LibreOffice 转换失败')
            shutil.copyfile(files[0], dest)
        return str(dest)
    if len(inputs) != 1: raise ValueError(f'「{op}」每次只能选择一个 PDF 文件。')
    d = fitz.open(inputs[0]); out = None
    try:
        if d.needs_pass and op != '解锁 PDF': raise ValueError('PDF 已加密，请先使用「解锁 PDF」。')
        if op == 'PDF 转文本':
            dest.write_text('\n\n'.join(f'--- 第 {i+1} 页 ---\n{p.get_text()}' for i,p in enumerate(d)), encoding='utf-8');return str(dest)
        if op == 'PDF 转 Word':
            from docx import Document
            doc = Document()
            for i,p in enumerate(d):
                if i:doc.add_page_break()
                text = p.get_text('text')
                for line in text.splitlines(): doc.add_paragraph(line)
            doc.save(str(dest));return str(dest)
        if op == 'PDF 转 Excel':
            from openpyxl import Workbook
            wb = Workbook();ws = wb.active;ws.title = '文本提取'
            ws.append(['页码','内容'])
            for i,p in enumerate(d):
                for s in p.get_text('text').splitlines():ws.append([i+1,s[:32767]])
            wb.save(str(dest));return str(dest)
        if op == 'PDF 转 PPT':
            from pptx import Presentation
            from pptx.util import Inches
            prs = Presentation();prs.slide_width=Inches(10);prs.slide_height=Inches(14)
            for page in d:
                slide=prs.slides.add_slide(prs.slide_layouts[6]);pix=page.get_pixmap(matrix=fitz.Matrix(1.3,1.3))
                bio=io.BytesIO(pix.tobytes('png'))
                slide.shapes.add_picture(bio,0,0,width=prs.slide_width,height=prs.slide_height)
            prs.save(str(dest));return str(dest)
        if op == 'PDF 转图片':
            # 多页统一 ZIP 打包，避免多下载窗口。
            with zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as z:
                for i,p in enumerate(d):
                    pix = p.get_pixmap(matrix=fitz.Matrix(dpi/72,dpi/72), alpha=False)
                    with Image.open(io.BytesIO(pix.tobytes('png'))) as im:
                        with io.BytesIO() as b:
                            im.convert('RGB').save(b, 'JPEG', quality=round(jpgquality*100), optimize=True)
                            z.writestr(f'page_{i+1:03}.jpg', b.getvalue())
            return str(dest)
        if op == '提取图片':
            count=0
            with zipfile.ZipFile(dest, 'w') as z:
                for i,p in enumerate(d):
                    for j,im in enumerate(p.get_images(full=True)):
                        dat = d.extract_image(im[0]);count +=1
                        z.writestr(f'page_{i+1:03}_{j+1:03}.{dat["ext"]}',dat['image'])
            if count==0:raise ValueError('未发现嵌入图片。')
            return str(dest)
        if op == '解锁 PDF':
            if not d.needs_pass: raise ValueError('该文件没有密码保护。')
            if not d.authenticate(extra): raise ValueError('密码不正确。')
            d.save(str(dest), encryption=fitz.PDF_ENCRYPT_NONE, garbage=4, deflate=True);return str(dest)
        if op == 'PDF 加密':
            if not extra: raise ValueError('请填写密码。')
            perm = fitz.PDF_PERM_PRINT|fitz.PDF_PERM_COPY
            d.save(str(dest), encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=extra, user_pw=extra, permissions=perm);return str(dest)
        if op == 'PDF 拆分（逐页）':
            nums=points(page_spec,len(d))
            with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
                for idx,n in enumerate(nums):
                    part=fitz.open();part.insert_pdf(d,from_page=n,to_page=n)
                    z.writestr(f'page_{n+1:03}_{idx+1:02}.pdf',part.tobytes(garbage=4,deflate=True));part.close()
            return str(dest)
        if op == '压缩 PDF（无损）':
            d.save(str(dest), garbage=4, deflate=True, clean=True);return str(dest)
        if op == '压缩 PDF（图像）':
            out=fitz.open()
            for p in d:
                pix=p.get_pixmap(matrix=fitz.Matrix(dpi/72,dpi/72),alpha=False)
                with Image.open(io.BytesIO(pix.tobytes('png'))) as im:
                    with io.BytesIO() as b:
                        im.convert('RGB').save(b,'JPEG',quality=round(jpgquality*100),optimize=True)
                        blob=b.getvalue()
                np=out.new_page(width=p.rect.width,height=p.rect.height)
                np.insert_image(np.rect,stream=blob)
            out.save(str(dest),garbage=4,deflate=True);return str(dest)
        out=fitz.open()
        if op in ('提取页面', '删除页面', '重新排列页面'):
            chosen=points(page_spec,len(d))
            selected = chosen if op!='删除页面' else [k for k in range(len(d)) if k not in set(chosen)]
            if not selected:raise ValueError('处理后没有剩余页面。')
            for n in selected:out.insert_pdf(d,from_page=n,to_page=n)
        else:
            out.insert_pdf(d)
            if op=='旋转 PDF':
                deg=int(extra or '90')
                for p in out:p.set_rotation((p.rotation+deg)%360)
            elif op=='添加页码':draw_text_overlay(out,'','页码')
            elif op=='添加文字水印':
                if not extra.strip():raise ValueError('请填写水印文本。')
                draw_text_overlay(out,extra,'水印')
            elif op=='删除 PDF 元数据':out.set_metadata({})
            else:raise ValueError(f'未实现的操作：{op}')
        out.save(str(dest),garbage=4,deflate=True);return str(dest)
    finally:
        d.close()
        if out:out.close()


TOOLS = [
 ('合并 PDF','PDF','pdf'),('PDF 拆分（逐页）','PDF','zip'),('提取页面','PDF','pdf'),('删除页面','PDF','pdf'),
 ('重新排列页面','PDF','pdf'),('旋转 PDF','PDF','pdf'),('添加页码','PDF','pdf'),('添加文字水印','PDF','pdf'),
 ('压缩 PDF（无损）','PDF','pdf'),('压缩 PDF（图像）','PDF','pdf'),('删除 PDF 元数据','PDF','pdf'),
 ('PDF 加密','安全','pdf'),('解锁 PDF','安全','pdf'),('PDF 转文本','转换','txt'),('PDF 转 Word','转换','docx'),
 ('PDF 转 Excel','转换','xlsx'),('PDF 转 PPT','转换','pptx'),('PDF 转图片','转换','zip'),
 ('提取图片','转换','zip'),('图片转 PDF','转换','pdf'),('Word 转 PDF','转换','pdf')]

def api_endpoint(base):
    """Chat Completions compatible endpoint only. Never POST to file/blob URLs."""
    from urllib.parse import urlsplit
    base=base.strip().rstrip('/')
    parsed=urlsplit(base)
    if (parsed.scheme not in ('http','https') or not parsed.hostname or parsed.query or parsed.fragment):
        raise ValueError('API 地址必须是服务商提供的 API Base URL，例如 https://api.openai.com/v1；不能填写网页、文件下载或含 ? 的签名链接。')
    if parsed.scheme!='https' and parsed.hostname not in ('localhost','127.0.0.1','::1'):
        raise ValueError('远程 API 必须使用 HTTPS；仅本机 localhost 可使用 HTTP。')
    if base.endswith('/chat/completions'): return base
    if base.endswith('/v1') or base.endswith('/v1/'): return base+'/chat/completions'
    # some OpenAI-compatible providers use /compatible-mode/v1 or /api/v1
    if '/v1' in parsed.path or '/compatible-mode/' in parsed.path:
        return base+'/chat/completions'
    raise ValueError('地址似乎不是兼容接口的 Base URL。一般应以 /v1 结尾；请到模型服务商的 API 文档获取地址。')

def chat_request(base,model,key,messages,timeout=45):
    url=api_endpoint(base)
    if not model.strip():raise ValueError('模型 ID 不能为空。')
    headers={'Content-Type':'application/json','Accept':'application/json','User-Agent':'oycy-PDF-Toolbox/2.0'}
    if key:headers['Authorization']='Bearer '+key
    payload={'model':model.strip(),'messages':messages}
    req=urllib.request.Request(url,data=json.dumps(payload,ensure_ascii=False).encode('utf-8'),headers=headers,method='POST')
    try:
        with urllib.request.urlopen(req,timeout=timeout) as response:
            result=json.loads(response.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        raw=e.read(900).decode('utf-8',errors='replace')
        suffix=(' 405 表示接口不支持 POST；常见原因：API 地址指向网页、对象存储或非兼容接口。请核对服务商的 Chat Completions Base URL。' if e.code==405 else
                ' 401/403 表示密钥无效或没有权限。' if e.code in (401,403) else
                ' 404 表示路径或模型不正确。' if e.code==404 else
                ' 429 表示请求额度或速率受限。' if e.code==429 else '')
        raise RuntimeError(f'HTTP {e.code}，请求地址：{url}\n{raw}{suffix}')
    except urllib.error.URLError as e:
        raise RuntimeError(f'无法连接 API：{e.reason}\n地址：{url}\n请检查网络、代理或 API 地址。')
    if not isinstance(result,dict) or not result.get('choices'):
        raise RuntimeError('接口返回成功，但没有 choices 字段；该服务可能不是 Chat Completions 兼容接口。')
    content=result['choices'][0].get('message',{}).get('content')
    if isinstance(content,list):content='\n'.join(p.get('text','') for p in content if isinstance(p,dict))
    return content or '接口返回了空内容。'

class App(tk.Tk):
    def __init__(self):
        super().__init__();enable_high_dpi(self);self.title(APP_NAME+' · WINDOWS EDITION')
        self.geometry('1220x820');self.minsize(940,700);self.configure(bg=PAPER)
        self.cfg=load_cfg();self.files=[];self.busy=False
        self.operation=tk.StringVar(value='合并 PDF');self.quality=tk.StringVar(value=self.cfg['quality'])
        self.page_spec=tk.StringVar();self.extra=tk.StringVar();self.output=tk.StringVar(value=self.cfg['output_dir'])
        self.api_base=tk.StringVar(value=self.cfg['api_base']);self.model=tk.StringVar(value=self.cfg['model']);self.key=tk.StringVar(value='')
        self.prompt=tk.StringVar(value='请用中文概括这份文件的要点，并列出重要数据。')
        self._style();self._header();self._footer();self._tabs()

    def _style(self):
        s=ttk.Style(self);s.theme_use('clam')
        s.configure('TNotebook',background=PAPER,borderwidth=0)
        s.configure('TNotebook.Tab',font=('Microsoft YaHei UI',11,'bold'),padding=(18,10),background='#e5dfd0',foreground=INK)
        s.map('TNotebook.Tab',background=[('selected',YELLOW)])
        s.configure('TLabel',background=PAPER,foreground=INK,font=('Microsoft YaHei UI',10))
        s.configure('TFrame',background=PAPER)
        s.configure('TEntry',padding=6)
        s.configure('TCombobox',padding=5)

    def _header(self):
        top=tk.Frame(self,bg=INK,height=96);top.pack(fill='x');top.pack_propagate(False)
        tk.Canvas(top,width=65,height=65,bg=INK,highlightthickness=0).place(x=20,y=15)
        c=tk.Canvas(top,width=65,height=65,bg=INK,highlightthickness=0);c.place(x=20,y=15)
        c.create_oval(2,2,62,62,fill=RED,outline='');c.create_rectangle(29,1,64,31,fill=YELLOW,outline='')
        tk.Label(top,text='oycy的PDF工具箱',bg=INK,fg='white',font=('Microsoft YaHei UI',25,'bold')).place(x=99,y=15)
        tk.Label(top,text='BAUHAUS / PDF STUDIO  ·  本地处理 + 自选 AI API',bg=INK,fg=YELLOW,font=('Arial',11,'bold')).place(x=102,y=64)

    def _tabs(self):
        tabs=ttk.Notebook(self);tabs.pack(fill='both',expand=True,padx=18,pady=15)
        tool=ttk.Frame(tabs);ai=ttk.Frame(tabs);settings=ttk.Frame(tabs)
        tabs.add(tool,text='01  PDF 工具');tabs.add(ai,text='02  AI 文档助手');tabs.add(settings,text='03  API 与设置')
        self._tools_ui(tool);self._ai_ui(ai);self._settings_ui(settings)

    def _label(self,parent,text,r,c=0,**kw):
        w=tk.Label(parent,text=text,bg=PAPER,fg=INK,font=('Microsoft YaHei UI',10,'bold'),anchor='w',**kw)
        w.grid(row=r,column=c,sticky='w',padx=10,pady=6);return w

    def _tools_ui(self,parent):
        parent.columnconfigure(1,weight=1)
        self._label(parent,'选择处理功能',0)
        choices=[x[0] for x in TOOLS]
        cb=ttk.Combobox(parent,textvariable=self.operation,values=choices,state='readonly',width=42)
        cb.grid(row=0,column=1,sticky='ew',padx=10,pady=8);cb.bind('<<ComboboxSelected>>',lambda _:self._hint())
        self._label(parent,'待处理文件',1)
        pane=tk.Frame(parent,bg=PAPER);pane.grid(row=1,column=1,sticky='ew',padx=10)
        self.file_label=tk.Label(pane,text='未选择文件',fg='#666',bg=PAPER,anchor='w',justify='left',wraplength=630)
        self.file_label.pack(fill='x')
        row=tk.Frame(parent,bg=PAPER);row.grid(row=2,column=1,sticky='w',padx=10,pady=(2,9))
        self._btn(row,'+ 添加文件',self.add_files,RED).pack(side='left',padx=(0,7))
        self._btn(row,'清空',self.clear_files,INK).pack(side='left',padx=(0,7))
        self._btn(row,'上移顺序',lambda:self._move(-1),BLUE).pack(side='left',padx=(0,7))
        self._btn(row,'下移顺序',lambda:self._move(1),BLUE).pack(side='left')
        self._label(parent,'输出质量',3)
        ttk.Combobox(parent,textvariable=self.quality,values=['轻量','标准','高清','超清'],state='readonly').grid(row=3,column=1,sticky='ew',padx=10,pady=7)
        self._label(parent,'页码 / 顺序',4)
        ttk.Entry(parent,textvariable=self.page_spec).grid(row=4,column=1,sticky='ew',padx=10,pady=7)
        self._label(parent,'附加输入',5)
        ttk.Entry(parent,textvariable=self.extra).grid(row=5,column=1,sticky='ew',padx=10,pady=7)
        self.hint=tk.Label(parent,text='',bg=PAPER,fg=BLUE,wraplength=770,justify='left',anchor='w')
        self.hint.grid(row=6,column=1,sticky='w',padx=10,pady=(4,10))
        self._btn(parent,'选择输出文件并开始处理',self.run_tool,RED,font_size=13).grid(row=7,column=1,sticky='ew',padx=10,pady=12)
        self._hint()
        tk.Label(parent,text='提示：页面范围填写 1,3,5-8；调整顺序可写 3,1,2。图像压缩会栅格化页面并失去文本可选性。',bg=PAPER,fg='#555',wraplength=760,justify='left').grid(row=8,column=1,sticky='w',padx=10,pady=6)

    def _hint(self):
        op=self.operation.get()
        hint={'旋转 PDF':'附加输入：旋转角度 90、180、270（默认 90）。','添加文字水印':'附加输入：填写水印文字，支持中文。',
              'PDF 加密':'附加输入：填写新密码。','解锁 PDF':'附加输入：填写已有密码。',
              'PDF 拆分（逐页）':'页码可留空表示全部页面。',
              '提取页面':'页码：如 1,3,5-8。','删除页面':'页码：填写要删除的页面。',
              '重新排列页面':'页码：如 3,1,2，必须明确所需页面顺序。',
              'PDF 转 Word':'提取文字生成 Word，不保证保留原始版式。',
              'PDF 转 Excel':'逐行提取文本到表格，不是高级表格识别。',
              'PDF 转 PPT':'每页成为一张图片幻灯片，不支持逐元素编辑。',
              'Word 转 PDF':'此工具需要另外安装 LibreOffice。'}
        self.hint.config(text=hint.get(op,'支持将结果保存到任意自选位置；根据工具功能选择 PDF 或图片。'))

    def _btn(self,parent,label,cmd,color,font_size=10):
        return tk.Button(parent,text=label,command=cmd,bg=color,fg='white' if color!=YELLOW else INK,
                         font=('Microsoft YaHei UI',font_size,'bold'),relief='flat',padx=13,pady=9,cursor='hand2',activebackground=YELLOW)

    def add_files(self):
        op=self.operation.get()
        if op=='图片转 PDF':types=[('图片文件','*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp')]
        elif op=='Word 转 PDF':types=[('Word 文件','*.docx *.doc')]
        else:types=[('PDF 文件','*.pdf')]
        files=filedialog.askopenfilenames(title='选择文件',filetypes=types+[('全部文件','*.*')])
        for f in files:
            if f not in self.files:self.files.append(f)
        self._refresh_files()

    def clear_files(self):self.files=[];self._refresh_files()
    def _move(self,n):
        if len(self.files)>1:
            self.files = self.files[1:]+self.files[:1] if n<0 else self.files[-1:]+self.files[:-1]
            self._refresh_files()
    def _refresh_files(self):
        txt='\n'.join(f'{i+1}. {Path(f).name}' for i,f in enumerate(self.files[:8]))
        if len(self.files)>8:txt+=f'\n...还有 {len(self.files)-8} 个文件'
        self.file_label.config(text=txt or '未选择文件')

    def _footer(self):
        self.status=tk.StringVar(value='就绪 · 本地 PDF 操作不上传文件')
        tk.Label(self,textvariable=self.status,bg=BLUE,fg='white',font=('Microsoft YaHei UI',10),anchor='w',padx=18,pady=8).pack(fill='x')
    def _async(self,job,done):
        if self.busy:return
        self.busy=True;self.status.set('● 正在执行，请勿关闭程序…');self.config(cursor='watch')
        def worker():
            try:ret=job(); self.after(0,lambda:finish(ret,None))
            except Exception as e:self.after(0,lambda err=str(e):finish(None,err))
        def finish(ret,err):
            self.busy=False;self.config(cursor='');self.status.set('✕ 处理失败，查看错误信息' if err else '✓ 处理完成')
            if err:messagebox.showerror('处理失败',err)
            else:done(ret)
        threading.Thread(target=worker,daemon=True).start()

    def run_tool(self):
        if not self.files:messagebox.showwarning('缺少文件','请先添加待处理文件');return
        op=self.operation.get();ext=next(x[2] for x in TOOLS if x[0]==op)
        name='oycy_'+re.sub(r'[^\w-]','_',op)+'.'+ext
        init=self.output.get() or str(Path.home()/'Downloads')
        path=filedialog.asksaveasfilename(title='选择输出文件保存位置',initialdir=init,initialfile=name,
               defaultextension='.'+ext,filetypes=[(ext.upper()+' 文件','*.'+ext),('全部文件','*.*')])
        if not path:return
        self._async(lambda:run_operation(op,self.files,path,self.quality.get(),self.page_spec.get(),self.extra.get()),
                    lambda ret:messagebox.showinfo('处理成功',f'已保存到：\n{ret}'))

    def _ai_ui(self,parent):
        parent.columnconfigure(0,weight=1)
        tk.Label(parent,text='AI 读取文档并生成分析',bg=PAPER,fg=INK,font=('Microsoft YaHei UI',17,'bold')).grid(row=0,column=0,sticky='w',padx=20,pady=(16,8))
        tk.Label(parent,text='点击选择 PDF / TXT / DOCX，提取文本后发送至你配置的 API 服务。请勿上传机密文件到不可信服务。',bg=PAPER,fg='#555',wraplength=810,justify='left').grid(row=1,column=0,sticky='w',padx=20)
        self.ai_file=tk.StringVar(value='尚未选择文档')
        self._btn(parent,'选择 AI 分析文件',self.pick_ai_file,BLUE).grid(row=2,column=0,sticky='w',padx=20,pady=9)
        tk.Label(parent,textvariable=self.ai_file,bg=PAPER,fg=INK,wraplength=820,justify='left').grid(row=3,column=0,sticky='w',padx=20)
        tk.Label(parent,text='输入你的分析问题：',bg=PAPER,fg=INK,font=('Microsoft YaHei UI',10,'bold')).grid(row=4,column=0,sticky='w',padx=20,pady=(12,3))
        self.ai_prompt=tk.Text(parent,height=4,wrap='word',font=('Microsoft YaHei UI',11))
        self.ai_prompt.grid(row=5,column=0,sticky='ew',padx=20)
        self.ai_prompt.insert('1.0',self.prompt.get())
        self._btn(parent,'使用所选模型分析文档',self.analyze,RED).grid(row=6,column=0,sticky='w',padx=20,pady=12)
        self.ai_result=tk.Text(parent,height=14,wrap='word',font=('Microsoft YaHei UI',10))
        self.ai_result.grid(row=7,column=0,sticky='nsew',padx=20,pady=(0,6));parent.rowconfigure(7,weight=1)
        self._btn(parent,'保存 AI 分析结果',self.save_ai_result,INK).grid(row=8,column=0,sticky='e',padx=20,pady=(0,10))
        self.ai_path=''

    def pick_ai_file(self):
        f=filedialog.askopenfilename(filetypes=[('支持的文档','*.pdf *.txt *.docx'),('全部文件','*.*')])
        if f:self.ai_path=f;self.ai_file.set(f)
    def read_ai_text(self,path):
        ext=Path(path).suffix.lower()
        if ext=='.txt':return Path(path).read_text(encoding='utf-8-sig',errors='replace')
        if ext=='.docx':
            from docx import Document
            return '\n'.join(p.text for p in Document(path).paragraphs)
        if ext=='.pdf':
            with fitz.open(path) as d:return '\n'.join(p.get_text() for p in d)
        raise ValueError('仅支持 PDF、TXT、DOCX')
    def _key(self):
        return self.key.get().strip() or keyring.get_password(SERVICE,'default') or ''
    def analyze(self):
        if not self.ai_path:messagebox.showwarning('缺少文件','请选择文档');return
        base=self.api_base.get().strip().rstrip('/')
        if not base.startswith('https://') and not base.startswith('http://localhost') and not base.startswith('http://127.0.0.1'):
            messagebox.showerror('不安全的 API 地址','请使用 HTTPS API 地址；本地模型可使用 localhost。');return
        key=self._key()
        if not key and 'localhost' not in base and '127.0.0.1' not in base:messagebox.showwarning('API Key','请先在 API 设置中填写并保存密钥。');return
        prompt=self.ai_prompt.get('1.0','end').strip()
        if not prompt:messagebox.showwarning('缺少问题','请输入分析问题');return
        def job():
            text=self.read_ai_text(self.ai_path)
            if not text.strip():raise ValueError('文档未提取到文字；扫描版 PDF 需要 OCR。')
            if len(text)>60000:text=text[:60000]+'\n[文件过长，已截取前 60000 字符]'
            messages=[{'role':'system','content':'你是严谨的中文文档分析助手。只根据给定文档作答，未知部分应说明。'},
                      {'role':'user','content':f'{prompt}\n\n---文档内容---\n{text}'}]
            return chat_request(base,self.model.get().strip(),key,messages,120)
        self._async(job,lambda ans:(self.ai_result.delete('1.0','end'),self.ai_result.insert('1.0',ans)))
    def save_ai_result(self):
        txt=self.ai_result.get('1.0','end').strip()
        if not txt:return
        f=filedialog.asksaveasfilename(defaultextension='.txt',filetypes=[('文本文件','*.txt')],initialfile='AI文档分析结果.txt')
        if f:Path(f).write_text(txt,encoding='utf-8');messagebox.showinfo('保存成功',f)

    def _settings_ui(self,parent):
        parent.columnconfigure(1,weight=1)
        tk.Label(parent,text='AI API 配置  /  自带密钥 BYOK',bg=PAPER,fg=INK,font=('Microsoft YaHei UI',17,'bold')).grid(row=0,column=0,columnspan=2,sticky='w',padx=20,pady=(22,17))
        self._label(parent,'API Base URL',1)
        ttk.Entry(parent,textvariable=self.api_base).grid(row=1,column=1,sticky='ew',padx=20,pady=8)
        self._label(parent,'模型 ID',2)
        ttk.Entry(parent,textvariable=self.model).grid(row=2,column=1,sticky='ew',padx=20,pady=8)
        self._label(parent,'API Key',3)
        ttk.Entry(parent,textvariable=self.key,show='•').grid(row=3,column=1,sticky='ew',padx=20,pady=8)
        tk.Label(parent,text='示例：OpenAI 兼容接口 /v1 + 模型名称。密钥保存在 Windows 凭据管理器，不写入设置 JSON。\n支持本地 OpenAI 兼容服务（例如 localhost:11434/v1）；不同服务商的接口可能存在兼容性差异。',
                 bg=PAPER,fg='#555',justify='left',wraplength=800).grid(row=4,column=0,columnspan=2,sticky='w',padx=20,pady=11)
        self._btn(parent,'保存 API 设置',self.save_settings,RED).grid(row=5,column=1,sticky='w',padx=20,pady=10)
        self._btn(parent,'测试 API 连接',self.test_api,BLUE).grid(row=5,column=0,sticky='w',padx=20,pady=10)
        self._btn(parent,'清除已保存的 API Key',self.delete_key,INK).grid(row=6,column=1,sticky='w',padx=20,pady=5)
        tk.Label(parent,text='隐私提示：基础 PDF 工具仅在本机处理。AI 分析会将提取出的文本上传到所配置的 API 端点；请确认服务提供商的数据政策。',bg=PAPER,fg=RED,wraplength=750,justify='left').grid(row=7,column=0,columnspan=2,sticky='w',padx=20,pady=20)
    def test_api(self):
        base=self.api_base.get().strip();model=self.model.get().strip();key=self._key()
        try:endpoint=api_endpoint(base)
        except Exception as e:messagebox.showerror('地址检查失败',str(e));return
        if not key and 'localhost' not in base and '127.0.0.1' not in base:
            messagebox.showwarning('缺少 API Key','请先填写 API Key（不必先保存），再测试连接。');return
        self.status.set('正在连接：'+endpoint+' ...')
        def job():return chat_request(base,model,key,[{'role':'user','content':'请只回复：连接成功'}],35)
        def done(result):
            self.status.set('API 测试成功')
            messagebox.showinfo('API 连接成功',f'接口：{endpoint}\n模型：{model}\n返回：{result[:350]}')
        self._async(job,done)

    def save_settings(self):
        base=self.api_base.get().strip()
        if not base.startswith('https://') and not base.startswith('http://localhost') and not base.startswith('http://127.0.0.1'):
            messagebox.showerror('API 地址错误','远程 API 请使用 HTTPS；本地可以使用 localhost。');return
        d=dict(self.cfg);d.update(api_base=base,model=self.model.get().strip(),quality=self.quality.get())
        save_cfg(d);self.cfg=d
        if self.key.get().strip():
            try:keyring.set_password(SERVICE,'default',self.key.get().strip());self.key.set('')
            except Exception as e:messagebox.showerror('密钥保存失败',str(e));return
        messagebox.showinfo('保存成功','API 地址和模型已保存；API Key 使用系统凭据存储。')
    def delete_key(self):
        try:keyring.delete_password(SERVICE,'default')
        except keyring.errors.PasswordDeleteError:pass
        self.key.set('');messagebox.showinfo('清除完成','已尝试清除 Windows 凭据中的 API Key。')


# ================ oycy Modern UI / soft Bauhaus ================
WHITE='#ffffff'; BG='#f5f7fb'; BORDER='#e5e9f0'; SOFT='#677488'; ACCENT='#1768ee'; GREEN='#2bbd86'
MODERN_TOOLS=[
 ('合并 PDF','合并PDF','将多个 PDF 合并为一个文件','↔','#ff5258','全部'),
 ('PDF 拆分（逐页）','分割PDF','按页码拆分并下载 ZIP','✂','#4388fa','全部'),
 ('压缩 PDF（无损）','无损压缩','优化文件结构，保留可选文字','⇣','#30bb85','全部'),
 ('PDF 转 Word','PDF转Word','提取文字为 DOCX（基础版）','W','#4585ee','转换'),
 ('PDF 转 PPT','PDF转PPT','每页转换为图片幻灯片','P','#ff8040','转换'),
 ('PDF 转 Excel','PDF转Excel','提取文本至电子表格','X','#32bf82','转换'),
 ('PDF 转图片','PDF转JPG','多页高清 JPG 打包 ZIP','▧','#ff8553','转换'),
 ('图片转 PDF','图片转PDF','将多张图片合并为 PDF','▣','#f95871','转换'),
 ('旋转 PDF','旋转PDF','批量旋转 PDF 页面','↻','#8a6cf5','页面'),
 ('添加页码','添加页码','为每页加入页码','12','#f5c33e','页面'),
 ('添加文字水印','文字水印','可输入中文水印文字','◈','#4c8bf8','页面'),
 ('重新排列页面','组织PDF','指定新页码顺序','▤','#8a93aa','页面'),
 ('提取页面','提取页面','选择需要保留的页面','▥','#446df1','页面'),
 ('删除页面','删除页面','移除指定 PDF 页面','−','#ee7778','页面'),
 ('PDF 加密','保护PDF','添加文件打开密码','♙','#ed5f88','安全'),
 ('解锁 PDF','解锁PDF','使用原密码解除保护','♧','#62baf6','安全'),
 ('PDF 转文本','PDF转文本','提取可选择的文本','T','#6d78ea','转换'),
 ('压缩 PDF（图像）','图像压缩','重绘页面以减小文件体积','▾','#27bca8','全部'),
 ('提取图片','提取图片','导出嵌入的所有图片','▧','#ff9564','转换'),
 ('Word 转 PDF','Word转PDF','需另装 LibreOffice','W','#518bdb','转换'),
 ('删除 PDF 元数据','清理元数据','删除文档元数据','◇','#7e8ca9','安全')]

class ModernApp(App):
    def __init__(self):
        self.filter_name='全部';self.search_text='';self.selected_idx=None
        self.card_views={}
        self.active_card_op=None
        super().__init__()
        self.title('oycy的PDF工具箱  ·  Desktop')
        self.geometry('1500x920');self.minsize(1180,760);self.configure(bg=BG)

    def _style(self):
        s=ttk.Style(self);s.theme_use('clam')
        s.configure('TNotebook',background=BG,borderwidth=0)
        s.configure('TNotebook.Tab',font=('Microsoft YaHei UI',12,'bold'),padding=(22,13),background=BG,foreground=SOFT,borderwidth=0)
        s.map('TNotebook.Tab',background=[('selected',WHITE)],foreground=[('selected',ACCENT)])
        s.configure('TFrame',background=BG)
        s.configure('TEntry',padding=10,font=('Microsoft YaHei UI',11))
        s.configure('TCombobox',padding=7)

    def _header(self):
        top=tk.Frame(self,bg=WHITE,height=68);top.pack(fill='x');top.pack_propagate(False)
        tk.Label(top,text='●',bg=WHITE,fg='#ff514b',font=('Arial',23)).pack(side='left',padx=(25,9))
        tk.Label(top,text='oycy的PDF工具箱',bg=WHITE,fg='#161c27',font=('Microsoft YaHei UI',18,'bold')).pack(side='left')
        tk.Label(top,text='  BAUHAUS / LOCAL  ',bg='#eaf2ff',fg=ACCENT,font=('Arial',9,'bold'),padx=8,pady=6).pack(side='left',padx=16)
        tk.Label(top,text='PDF · 文档转换 · AI',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',10)).pack(side='right',padx=28)
        tk.Frame(self,bg=BORDER,height=1).pack(fill='x')

    def _tabs(self):
        tabs=ttk.Notebook(self);tabs.pack(fill='both',expand=True,padx=16,pady=(8,4))
        home=ttk.Frame(tabs);ai=ttk.Frame(tabs);settings=ttk.Frame(tabs)
        tabs.add(home,text='  全部 PDF 工具  ');tabs.add(ai,text='  ✦ AI 文档助手  ');tabs.add(settings,text='  ⚙ API 与设置  ')
        self.tabs=tabs
        self._tools_ui(home)
        self._ai_ui(ai);self._settings_ui(settings)

    def _panel(self,parent,padx=14,pady=14):
        f=tk.Frame(parent,bg=WHITE,highlightbackground=BORDER,highlightthickness=1)
        return f

    def _button(self,parent,text,cmd,primary=False):
        return tk.Button(parent,text=text,command=cmd,bg=ACCENT if primary else WHITE,
            fg=WHITE if primary else '#263246',activebackground='#105bd4' if primary else '#eaf2ff',
            activeforeground=WHITE if primary else ACCENT,relief='flat',bd=0,
            highlightthickness=1,highlightbackground=ACCENT if primary else BORDER,
            font=('Microsoft YaHei UI',10,'bold'),cursor='hand2',padx=13,pady=9)

    def _tools_ui(self,parent):
        parent.grid_columnconfigure(0,weight=1);parent.grid_rowconfigure(0,weight=1)
        self.horizontal_pane=tk.PanedWindow(parent,orient='horizontal',bg=BORDER,bd=0,sashwidth=9,sashrelief='flat',opaqueresize=True)
        self.horizontal_pane.grid(row=0,column=0,sticky='nsew',padx=3,pady=2)
        main=tk.PanedWindow(self.horizontal_pane,orient='vertical',bg=BORDER,bd=0,sashwidth=9,opaqueresize=True)
        side=tk.Frame(self.horizontal_pane,bg=BG,width=285)
        self.horizontal_pane.add(main,minsize=560,stretch='always');self.horizontal_pane.add(side,minsize=235,stretch='never')
        upper=tk.Frame(main,bg=BG);lower=tk.Frame(main,bg=BG)
        main.add(upper,minsize=190,stretch='always');main.add(lower,minsize=145,stretch='never')
        self.vertical_pane=main
        self.after(250,lambda:self._init_sashes())

        hero=tk.Frame(upper,bg=BG);hero.pack(fill='x',pady=(10,6))
        tk.Label(hero,text='所有处理 PDF 所需的工具都集中在一个地方',bg=BG,fg='#141b27',font=('Microsoft YaHei UI',21,'bold'),anchor='w').pack(fill='x',padx=6)
        tk.Label(hero,text='简单易用  ·  本地处理  ·  安全高效    选择工具后添加文件，设置参数并导出结果。',bg=BG,fg=SOFT,font=('Microsoft YaHei UI',10),anchor='w').pack(fill='x',padx=8,pady=(6,6))
        bar=tk.Frame(upper,bg=BG);bar.pack(fill='x',pady=(4,12))
        self.search_var=tk.StringVar()
        search=ttk.Entry(bar,textvariable=self.search_var,width=26);search.pack(side='left',padx=(5,8),ipady=4)
        search.insert(0,'')
        self.search_var.trace_add('write',lambda *a:self._render_cards())
        for t in ('全部','转换','页面','安全'):
            self._button(bar,t,lambda v=t:self._set_filter(v)).pack(side='left',padx=3)
        self.card_holder=self._panel(upper);self.card_holder.pack(fill='both',expand=True,padx=3,pady=(0,8))
        # scrollable tool cards
        self.tools_canvas=tk.Canvas(self.card_holder,bg=WHITE,highlightthickness=0)
        self.tools_scroll=ttk.Scrollbar(self.card_holder,orient='vertical',command=self.tools_canvas.yview)
        self.tools_canvas.configure(yscrollcommand=self.tools_scroll.set)
        self.tools_scroll.pack(side='right',fill='y');self.tools_canvas.pack(side='left',fill='both',expand=True)
        self.inner=tk.Frame(self.tools_canvas,bg=WHITE)
        self.canvas_window=self.tools_canvas.create_window((0,0),window=self.inner,anchor='nw')
        self.inner.bind('<Configure>',lambda e:self.tools_canvas.configure(scrollregion=self.tools_canvas.bbox('all')))
        self.tools_canvas.bind('<Configure>',lambda e:self.tools_canvas.itemconfigure(self.canvas_window,width=e.width))
        self.tools_canvas.bind_all('<MouseWheel>',self._wheel)
        self._render_cards()
        upload=self._panel(side);upload.pack(fill='x',pady=(4,10))
        tk.Label(upload,text='☁',bg=WHITE,fg=ACCENT,font=('Arial',27,'bold')).pack(pady=(13,2))
        tk.Label(upload,text='添加待处理文件',bg=WHITE,fg=ACCENT,font=('Microsoft YaHei UI',12,'bold')).pack()
        tk.Label(upload,text='支持一次选择多个文件',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9)).pack(pady=5)
        self._button(upload,'＋ 选择文件',self.add_files,True).pack(fill='x',padx=17,pady=(9,16))
        options=self._panel(side);options.pack(fill='x',pady=(0,10))
        tk.Label(options,text='输出质量',bg=WHITE,fg='#17243a',font=('Microsoft YaHei UI',12,'bold'),anchor='w').pack(fill='x',padx=16,pady=(15,7))
        for val,desc in (('轻量','体积优先'),('标准','质量与体积平衡'),('高清','清晰度优先'),('超清','细节优先')):
            row=tk.Frame(options,bg=WHITE);row.pack(fill='x',padx=12,pady=1)
            tk.Radiobutton(row,text=f'{val}    · {desc}',variable=self.quality,value=val,bg=WHITE,fg='#344057',selectcolor=WHITE,activebackground=WHITE,font=('Microsoft YaHei UI',9)).pack(anchor='w')
        tk.Label(options,text='部分无损操作不会改变画质',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',8)).pack(anchor='w',padx=17,pady=(6,14))
        settings=self._panel(side);settings.pack(fill='x',pady=(0,10))
        tk.Label(settings,text='当前工具 / 输出',bg=WHITE,fg='#17243a',font=('Microsoft YaHei UI',11,'bold'),anchor='w').pack(fill='x',padx=15,pady=(12,5))
        self.op_label=tk.Label(settings,text='合并PDF',bg=WHITE,fg=ACCENT,font=('Microsoft YaHei UI',12,'bold'),anchor='w');self.op_label.pack(fill='x',padx=15)
        self.file_label=tk.Label(settings,text='尚未添加文件',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9),anchor='w',justify='left',wraplength=230)
        self.file_label.pack(fill='x',padx=15,pady=7)
        self._button(settings,'开始处理 · 选择保存位置',self.run_tool,True).pack(fill='x',padx=12,pady=(6,12))
        helpbox=self._panel(side);helpbox.pack(fill='x')
        tk.Label(helpbox,text='✦ AI 文档助手',bg=WHITE,fg='#7752d9',font=('Microsoft YaHei UI',11,'bold'),anchor='w').pack(fill='x',padx=14,pady=(12,5))
        tk.Label(helpbox,text='文档总结、内容问答、自定义 API',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9)).pack(anchor='w',padx=14)
        self._button(helpbox,'打开 AI 助手 →',lambda:self.tabs.select(1)).pack(fill='x',padx=12,pady=(8,12))
        self._workflow(lower)
        self._select_tool(self.operation.get())
        self._hint()

    def _init_sashes(self):
        try:
            w=self.horizontal_pane.winfo_width();h=self.vertical_pane.winfo_height()
            if w>850:self.horizontal_pane.sash_place(0,max(560,w-290),0)
            if h>500:self.vertical_pane.sash_place(0,0,max(260,int(h*.65)))
            self.status.set('就绪 · 可拖动工作区之间的分隔线调整大小')
        except tk.TclError:pass

    def _wheel(self,e):
        if self.tabs.index('current')==0 and self.tools_canvas.winfo_exists():
            self.tools_canvas.yview_scroll(int(-e.delta/120),'units')
    def _set_filter(self,value):
        self.filter_name=value;self._render_cards()

    def _preview_card(self, op, hovering):
        if op == self.active_card_op:
            return
        view=self.card_views.get(op)
        if not view:
            return
        border = '#c9d5ea' if hovering else BORDER
        bg = '#fbfdff' if hovering else WHITE
        view['card'].configure(highlightbackground=border, bg=bg)
        view['txt'].configure(bg=bg)
        for key in ('title','desc','arrow'):
            try:view[key].configure(bg=bg)
            except Exception:pass
        if hovering:
            try:view['arrow'].configure(fg=ACCENT)
            except Exception:pass
        else:
            try:view['arrow'].configure(fg='#7f8da5')
            except Exception:pass

    def _refresh_card_highlight(self):
        for op,view in getattr(self,'card_views',{}).items():
            selected = (op == self.active_card_op)
            bg = '#f4f8ff' if selected else WHITE
            border = ACCENT if selected else BORDER
            view['card'].configure(bg=bg, highlightbackground=border, highlightcolor=ACCENT)
            view['txt'].configure(bg=bg)
            view['title'].configure(bg=bg, fg=ACCENT if selected else '#152033')
            view['desc'].configure(bg=bg, fg='#4e5d75' if selected else SOFT)
            view['arrow'].configure(bg=bg, fg=ACCENT if selected else '#7f8da5')

    def _render_cards(self):
        if not hasattr(self,'inner'):return
        for w in self.inner.winfo_children():w.destroy()
        self.card_views={}
        q=self.search_var.get().strip().lower()
        items=[it for it in MODERN_TOOLS if (self.filter_name=='全部' or it[5]==self.filter_name) and (not q or q in (it[1]+' '+it[0]+' '+it[2]).lower())]
        for j,it in enumerate(items):
            op,title,desc,icon,color,group=it;r,c=divmod(j,3)
            self.inner.grid_columnconfigure(c,weight=1,uniform='cards')
            card=tk.Frame(self.inner,bg=WHITE,highlightthickness=2,highlightbackground=BORDER,highlightcolor=ACCENT,cursor='hand2')
            card.grid(row=r,column=c,sticky='nsew',padx=8,pady=8,ipady=2)
            left=tk.Frame(card,bg=color,width=52,height=52);left.pack(side='left',padx=(14,12),pady=16);left.pack_propagate(False)
            icon_label=tk.Label(left,text=icon,bg=color,fg='white',font=('Arial',20,'bold'))
            icon_label.pack(expand=True)
            txt=tk.Frame(card,bg=WHITE);txt.pack(side='left',fill='both',expand=True,pady=13)
            title_lbl=tk.Label(txt,text=title,bg=WHITE,fg='#152033',font=('Microsoft YaHei UI',12,'bold'),anchor='w')
            title_lbl.pack(fill='x')
            desc_lbl=tk.Label(txt,text=desc,bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9),anchor='w',wraplength=220,justify='left')
            desc_lbl.pack(fill='x',pady=(4,0))
            arrow_lbl=tk.Label(card,text='›',bg=WHITE,fg='#7f8da5',font=('Arial',24))
            arrow_lbl.pack(side='right',padx=12)
            self.card_views[op]={'card':card,'txt':txt,'title':title_lbl,'desc':desc_lbl,'arrow':arrow_lbl,'left':left,'icon':icon_label}
            def choose(e=None,o=op): self._select_tool(o)
            def enter(e=None,o=op): self._preview_card(o,True)
            def leave(e=None,o=op): self._preview_card(o,False)
            binders=[card,left,txt,icon_label,title_lbl,desc_lbl,arrow_lbl]
            for w in binders:
                w.bind('<Button-1>',choose)
                w.bind('<Enter>',enter)
                w.bind('<Leave>',leave)
        self._refresh_card_highlight()
        if not items:tk.Label(self.inner,text='未找到匹配的工具',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',12)).pack(pady=45)

    def _workflow(self,parent):
        block=self._panel(parent);block.pack(fill='both',expand=True,padx=3,pady=(0,3))
        tk.Label(block,text='处理工作区',bg=WHITE,fg='#16213a',font=('Microsoft YaHei UI',12,'bold')).pack(anchor='w',padx=16,pady=(10,3))
        row=tk.Frame(block,bg=WHITE);row.pack(fill='both',expand=True,padx=12,pady=4)
        self.files_list=tk.Listbox(row,height=5,selectmode='single',bg='#f9fbff',fg='#223044',font=('Microsoft YaHei UI',11),bd=0,highlightthickness=1,highlightbackground=BORDER,activestyle='none')
        self.files_list.pack(side='left',fill='both',expand=True,padx=(4,12));self.files_list.bind('<<ListboxSelect>>',self._on_selection)
        actions=tk.Frame(row,bg=WHITE);actions.pack(side='left')
        for txt,callback in (('上移',lambda:self._reorder(-1)),('下移',lambda:self._reorder(1)),('删除',self._remove_selected),('清空',self.clear_files)):
            self._button(actions,txt,callback).pack(fill='x',pady=2)
        form=tk.Frame(block,bg=WHITE);form.pack(fill='x',padx=17,pady=(6,9))
        tk.Label(form,text='页码 / 顺序',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9)).grid(row=0,column=0,sticky='w')
        ttk.Entry(form,textvariable=self.page_spec,width=22).grid(row=0,column=1,sticky='ew',padx=(8,18))
        tk.Label(form,text='水印 / 密码 / 角度',bg=WHITE,fg=SOFT,font=('Microsoft YaHei UI',9)).grid(row=0,column=2,sticky='w')
        ttk.Entry(form,textvariable=self.extra,width=22).grid(row=0,column=3,sticky='ew',padx=(8,0))
        form.grid_columnconfigure(1,weight=1);form.grid_columnconfigure(3,weight=1)
        self.hint=tk.Label(block,text='',bg=WHITE,fg=SOFT,wraplength=850,justify='left',anchor='w',font=('Microsoft YaHei UI',9))
        self.hint.pack(fill='x',padx=16,pady=(0,12))

    def _select_tool(self,operation):
        self.operation.set(operation)
        self.active_card_op=operation
        title=next((x[1] for x in MODERN_TOOLS if x[0]==operation),operation)
        self.op_label.config(text=title)
        if hasattr(self,'status'):
            self.status.set('已选择工具：'+title+' · 卡片已高亮，可继续添加文件并开始处理')
        self._refresh_card_highlight()
        self._hint()
    def _hint(self):
        op=self.operation.get()
        tips={'PDF 转 Word':'基础版：文字提取为 Word，不保留复杂排版。','PDF 转 Excel':'基础版：按文本行导出，并非表格结构识别。','PDF 转 PPT':'每页以图片加入 PPT，页面内容不可逐项编辑。',
            '压缩 PDF（图像）':'图像模式将栅格化页面，可能丢失文本可选性。','压缩 PDF（无损）':'无损优化不保证文件大小一定减小。','Word 转 PDF':'需要在电脑上另外安装 LibreOffice。',
            '重新排列页面':'在「页码 / 顺序」输入新的页序，例如 3,1,2。','PDF 拆分（逐页）':'留空拆分全部页面；输入 1,3,5-8 指定范围。',
            '添加文字水印':'请在「水印 / 密码 / 角度」填写中文或英文水印。','旋转 PDF':'可输入 90、180、270，默认 90。','PDF 加密':'在「水印 / 密码 / 角度」填写新密码。','解锁 PDF':'请输入现有密码。'}
        self.hint.config(text='提示 · '+tips.get(op,'选择文件，必要时调整页面范围或附加输入，点击右侧按钮即可选择保存位置。'))
    def _on_selection(self,e=None):
        v=self.files_list.curselection();self.selected_idx=v[0] if v else None
    def _reorder(self,delta):
        idx=self.selected_idx
        if idx is None:return
        dst=idx+delta
        if 0<=dst<len(self.files):
            self.files[idx],self.files[dst]=self.files[dst],self.files[idx];self._refresh_files()
            self.files_list.selection_set(dst);self.selected_idx=dst
    def _remove_selected(self):
        if self.selected_idx is not None and self.selected_idx<len(self.files):
            del self.files[self.selected_idx];self.selected_idx=None;self._refresh_files()
    def _refresh_files(self):
        if not hasattr(self,'files_list'):return
        self.files_list.delete(0,'end')
        for i,f in enumerate(self.files):self.files_list.insert('end',f'{i+1:02d}   {Path(f).name}')
        self.file_label.config(text=f'已添加 {len(self.files)} 个文件' if self.files else '尚未添加文件')
        if hasattr(self,'status'):self.status.set(f'文件列表已更新：{len(self.files)} 个文件')
    def _footer(self):
        self.status=tk.StringVar(value='就绪  ·  本地处理文件，AI 分析时需自行选择可信 API')
        tk.Label(self,textvariable=self.status,bg='#eaf2ff',fg=ACCENT,font=('Microsoft YaHei UI',10,'bold'),anchor='w',padx=17,pady=9).pack(fill='x')

if __name__=='__main__':
    app = ModernApp()
    if '--smoke-test' in sys.argv:
        app.update_idletasks()
        assert len(app.card_views) >= 15, '工具卡片未正确加载'
        app._select_tool('PDF 转 PPT')
        assert app.active_card_op == 'PDF 转 PPT'
        print('GUI smoke test passed')
        app.destroy()
    else:
        app.mainloop()
